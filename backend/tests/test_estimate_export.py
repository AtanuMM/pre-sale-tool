from __future__ import annotations

import csv
from io import BytesIO, StringIO
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.documents.csv_export import csv_safe_cell
from tests.test_estimate import (
    _approve_through_features,
    _gemini_batch,
    _sample_planning,
    _team_inputs,
)
from tests.test_steps import _auth, _create_project_multipart

pytest_plugins = ["tests.test_steps"]


def _inject_title() -> str:
    return '=HYPERLINK("http://x","y")'


@patch("app.services.step_generation.generate_structured_result")
def test_estimate_exports_docx_xlsx_csv(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    _, settings = seeded_session
    headers = _auth(client, settings)
    pid = _create_project_multipart(client, headers)["project"]["id"]
    _approve_through_features(client, headers, mock_gen, pid)

    planning = _sample_planning().model_dump()
    planning["exclusions"] = [_inject_title(), "=1+1"]
    from app.llm.gemini import StructuredGenerationResult
    from app.steps.schemas.estimate_v2 import EstimatePlanningOutput

    mock_gen.side_effect = [
        StructuredGenerationResult(
            data=EstimatePlanningOutput.model_validate(planning),
            model_id="m",
            tokens_in=1,
            tokens_out=1,
            api_attempt_count=1,
        ),
        _gemini_batch(),
    ]
    vid = client.post(
        f"/projects/{pid}/steps/estimate/generate",
        headers=headers,
        json={"inputs": _team_inputs()},
    ).json()["version"]["id"]

    bad = client.get(f"/step-versions/{vid}/download?format=pdf", headers=headers)
    assert bad.status_code == 422

    docx_resp = client.get(f"/step-versions/{vid}/download?format=docx", headers=headers)
    assert docx_resp.status_code == 200
    assert "DRAFT" in docx_resp.headers.get("content-disposition", "")
    doc = Document(BytesIO(docx_resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Estimate health" in text or "Summary" in text or "Recommended" in text

    xlsx_resp = client.get(f"/step-versions/{vid}/download?format=xlsx", headers=headers)
    assert xlsx_resp.status_code == 200
    wb = load_workbook(BytesIO(xlsx_resp.content))
    assert set(wb.sheetnames) >= {
        "Summary",
        "Team",
        "Settings",
        "Estimate",
        "Phases",
        "Components",
        "Assumptions",
        "Exclusions",
    }
    estimate = wb["Estimate"]
    formulas = [
        cell.value
        for row in estimate.iter_rows(min_row=2)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert formulas
    assert any("CapHigh" in f or "MIN" in f or "BufferHigh" in f for f in formulas)
    exclusions = wb["Exclusions"]
    found_injection = False
    for row in exclusions.iter_rows(min_row=2, max_col=1):
        cell = row[0]
        if cell.value and str(cell.value).startswith("="):
            assert cell.data_type == "s"
            found_injection = True
    assert found_injection

    csv_resp = client.get(f"/step-versions/{vid}/download?format=csv", headers=headers)
    assert csv_resp.status_code == 200
    raw = csv_resp.content
    assert raw.startswith(b"\xef\xbb\xbf")
    text_csv = raw.decode("utf-8-sig")
    assert "\r\n" in text_csv
    reader = csv.reader(StringIO(text_csv))
    header = next(reader)
    assert header[0] == "Status"
    first = next(reader)
    assert first[3].startswith("F-")


def test_legacy_v1_docx_layout() -> None:
    from app.documents.layouts import layout_for_step
    from app.documents.meta import ExportMeta
    from app.documents.render import render_step_docx
    from app.steps.layouts_registry import ESTIMATE_LAYOUT

    content = {
        "summary": "Old",
        "approach": "API",
        "components": [],
        "phases": [],
        "assumptions_and_solutions": [],
        "exclusions": [],
        "feature_estimates": [],
        "work_packages": [],
        "effort_lines": [],
        "computed": {
            "cards": [{"key": "hours_with_ai", "label": "Hours with AI tools", "value": "10 h"}],
            "method_note": "v1 method",
        },
    }
    layout = layout_for_step("estimate", content)
    assert layout == ESTIMATE_LAYOUT
    meta = ExportMeta(
        step_key="estimate",
        step_title="Estimate",
        project_name="P",
        client_name="C",
        version_no=1,
        status="in_review",
        status_label="in review",
        marker_line="DRAFT — not approved",
        metadata_date_label="Generated",
        metadata_date_value="1 January 2026",
        created_by_name="Author",
        is_draft_filename=True,
    )
    blob = render_step_docx("estimate", content, meta)
    assert blob.startswith(b"PK")


def test_csv_injection_prefix() -> None:
    assert csv_safe_cell('=HYPERLINK("http://x","y")').startswith("'")
    assert csv_safe_cell("=1+1").startswith("'")
    assert csv_safe_cell("+cmd").startswith("'")
    assert csv_safe_cell("normal") == "normal"
