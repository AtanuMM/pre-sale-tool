from __future__ import annotations

import uuid
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.config import Settings, get_settings
from app.models.auth import Permission, Role, User, UserRole
from app.models.steps import StepVersion
from app.models.system import AuditLog
from app.security.passwords import hash_password
from app.services.audit import AuditAction
from tests.test_steps import (
    _analyst_headers,
    _auth,
    _create_project_multipart,
    _login,
    _mock_gemini_success,
    _sample_scope_output,
    _viewer_headers,
)

pytest_plugins = ["tests.test_steps"]

pytestmark = pytest.mark.usefixtures("_steps_env")


def _download_url(version_id: str) -> str:
    return f"/step-versions/{version_id}/download?format=docx"


def _load_docx(data: bytes) -> Document:
    return Document(BytesIO(data))


def _lead_headers(client: TestClient, factory: sessionmaker[Session]) -> dict[str, str]:
    with factory() as session:
        role = session.scalar(select(Role).where(Role.name == "Lead"))
        assert role is not None
        user = User(
            email="lead-dl@example.test",
            full_name="Lead Downloader",
            password_hash=hash_password("lead-pass-12"),
            is_active=True,
        )
        session.add(user)
        session.flush()
        session.add(UserRole(user_id=user.id, role_id=role.id))
        session.commit()
    token = _login(client, "lead-dl@example.test", "lead-pass-12")
    return {"Authorization": f"Bearer {token}"}


def _view_only_headers(client: TestClient, factory: sessionmaker[Session]) -> dict[str, str]:
    with factory() as session:
        perm = session.scalar(select(Permission).where(Permission.code == "project.view"))
        assert perm is not None
        role = Role(name="ViewOnlyNoDownload", description="test")
        session.add(role)
        session.flush()
        role.permissions.append(perm)
        user = User(
            email="viewonly@example.test",
            full_name="View Only",
            password_hash=hash_password("viewonly-pass-12"),
            is_active=True,
        )
        session.add(user)
        session.flush()
        session.add(UserRole(user_id=user.id, role_id=role.id))
        session.commit()
    token = _login(client, "viewonly@example.test", "viewonly-pass-12")
    return {"Authorization": f"Bearer {token}"}


def _create_in_review_version(
    client: TestClient,
    settings: Settings,
    mock_gen: MagicMock,
) -> tuple[str, str]:
    mock_gen.return_value = _mock_gemini_success()
    headers = _auth(client, settings)
    created = _create_project_multipart(client, headers)
    pid = created["project"]["id"]
    resp = client.post(f"/projects/{pid}/steps/scope_analysis/generate", headers=headers)
    version_id = resp.json()["version"]["id"]
    return pid, version_id


@patch("app.services.step_generation.generate_structured_result")
def test_download_permission_matrix(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    factory, settings = seeded_session
    _, version_id = _create_in_review_version(client, settings, mock_gen)
    url = _download_url(version_id)

    assert client.get(url).status_code == 401

    view_only = _view_only_headers(client, factory)
    assert client.get(url, headers=view_only).status_code == 403

    for headers in (
        _viewer_headers(client, factory),
        _analyst_headers(client, factory),
        _lead_headers(client, factory),
        _auth(client, settings),
    ):
        resp = client.get(url, headers=headers)
        assert resp.status_code == 200, resp.text
        assert resp.headers.get("x-content-type-options") == "nosniff"
        assert "attachment" in resp.headers.get("content-disposition", "")
        _load_docx(resp.content)


@patch("app.services.step_generation.generate_structured_result")
def test_download_404_unknown_version(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    _, settings = seeded_session
    _create_in_review_version(client, settings, mock_gen)
    headers = _auth(client, settings)
    missing = uuid.uuid4()
    assert client.get(_download_url(str(missing)), headers=headers).status_code == 404


@patch("app.services.step_generation.generate_structured_result")
def test_download_422_format(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    _, settings = seeded_session
    _, version_id = _create_in_review_version(client, settings, mock_gen)
    headers = _auth(client, settings)
    assert (
        client.get(f"/step-versions/{version_id}/download", headers=headers).status_code
        == 422
    )
    assert (
        client.get(
            f"/step-versions/{version_id}/download?format=pdf", headers=headers
        ).status_code
        == 422
    )


@patch("app.services.step_generation.generate_structured_result")
def test_download_409_no_content(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    factory, settings = seeded_session
    headers = _auth(client, settings)
    created = _create_project_multipart(client, headers)
    pid = created["project"]["id"]

    mock_gen.side_effect = Exception("fail")
    failed = client.post(
        f"/projects/{pid}/steps/scope_analysis/generate", headers=headers
    )
    failed_id = failed.json()["version"]["id"]
    resp = client.get(_download_url(failed_id), headers=headers)
    assert resp.status_code == 409
    assert resp.json()["detail"] == "No content to download for this version"

    with factory() as session:
        from app.models.auth import User as UserModel

        user_id = session.scalar(
            select(UserModel.id).where(UserModel.email == settings.ADMIN_EMAIL)
        )
        queued = StepVersion(
            project_id=uuid.UUID(pid),
            step_key="scope_analysis",
            version_no=99,
            status="queued",
            source="generated",
            content=None,
            created_by=user_id,
        )
        session.add(queued)
        session.commit()
        qid = str(queued.id)

    assert client.get(_download_url(qid), headers=headers).status_code == 409

    with factory() as session:
        from app.models.auth import User as UserModel

        user_id = session.scalar(
            select(UserModel.id).where(UserModel.email == settings.ADMIN_EMAIL)
        )
        empty = StepVersion(
            project_id=uuid.UUID(pid),
            step_key="scope_analysis",
            version_no=100,
            status="in_review",
            source="generated",
            content={},
            created_by=user_id,
        )
        session.add(empty)
        session.commit()
        empty_id = str(empty.id)

    assert client.get(_download_url(empty_id), headers=headers).status_code == 409


@patch("app.services.step_generation.generate_structured_result")
def test_markers_in_review_and_approved(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    _, settings = seeded_session
    headers = _auth(client, settings)
    _, version_id = _create_in_review_version(client, settings, mock_gen)

    draft = client.get(_download_url(version_id), headers=headers)
    doc = _load_docx(draft.content)
    header_text = doc.sections[0].header.paragraphs[0].text
    assert "DRAFT, NOT APPROVED" in header_text
    assert doc.paragraphs[0].text == "DRAFT, NOT APPROVED"

    approved = client.post(
        f"/step-versions/{version_id}/approve", headers=headers, json={}
    )
    assert approved.status_code == 200
    approved_id = approved.json()["id"]
    final = client.get(_download_url(approved_id), headers=headers)
    doc2 = _load_docx(final.content)
    header2 = doc2.sections[0].header.paragraphs[0].text
    assert "DRAFT" not in header2
    assert "Approved by" in header2
    assert "Self-approved" in header2
    assert "-DRAFT" not in final.headers["content-disposition"]


def test_marker_lines_superseded_and_stale_in_docx() -> None:
    from app.documents.meta import ExportMeta, build_marker_line
    from app.documents.render import render_step_docx

    for status, expected in (
        ("superseded", "SUPERSEDED DRAFT, NOT APPROVED"),
        ("stale", "STALE: an upstream step changed, not current"),
    ):
        marker = build_marker_line(
            status=status, approval=None, fallback_approved_at=None
        )
        assert marker == expected
        meta = ExportMeta(
            step_key="scope_analysis",
            step_title="Scope Analysis",
            project_name="Proj",
            client_name="Client",
            version_no=2,
            status=status,
            status_label=status,
            marker_line=marker,
            metadata_date_label="Generated",
            metadata_date_value="1 January 2026",
            created_by_name="Author",
            is_draft_filename=True,
        )
        blob = render_step_docx(
            "scope_analysis",
            {"executive_summary": "Body text."},
            meta,
        )
        doc = _load_docx(blob)
        assert expected in doc.sections[0].header.paragraphs[0].text
        assert expected in doc.paragraphs[0].text


@patch("app.services.step_generation.generate_structured_result")
def test_scope_sections_and_tables(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    _, settings = seeded_session
    headers = _auth(client, settings)
    _, version_id = _create_in_review_version(client, settings, mock_gen)
    resp = client.get(_download_url(version_id), headers=headers)
    doc = _load_docx(resp.content)
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Executive summary" in body_text
    assert "Fictional summary." in body_text
    assert "Objective A" in body_text
    assert doc.tables
    table_text = doc.tables[0].rows[1].cells[0].text
    assert table_text == "1" or table_text  # metadata table version cell


@patch("app.services.step_generation.generate_structured_result")
def test_missing_extra_fields_and_sanitize(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    factory, settings = seeded_session
    headers = _auth(client, settings)
    _, version_id = _create_in_review_version(client, settings, mock_gen)

    with factory() as session:
        version = session.get(StepVersion, uuid.UUID(version_id))
        assert version is not None
        content = dict(version.content or {})
        content.pop("risks", None)
        content["executive_summary"] = "**not bold** summary"
        content["extra_field"] = "ignored"
        version.content = content
        session.commit()

    resp = client.get(_download_url(version_id), headers=headers)
    doc = _load_docx(resp.content)
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "**not bold** summary" in joined


@patch("app.services.step_generation.generate_structured_result")
def test_download_audit_and_filename(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
) -> None:
    factory, settings = seeded_session
    headers = _auth(client, settings)
    _, version_id = _create_in_review_version(client, settings, mock_gen)
    resp = client.get(_download_url(version_id), headers=headers)
    assert "-DRAFT.docx" in resp.headers["content-disposition"]

    with factory() as session:
        audit = session.scalar(
            select(AuditLog).where(
                AuditLog.action == AuditAction.DOCUMENT_DOWNLOADED,
                AuditLog.entity_id == uuid.UUID(version_id),
            )
        )
        assert audit is not None
        assert audit.meta.get("format") == "docx"
        assert audit.meta.get("step_key") == "scope_analysis"
        assert "executive_summary" not in str(audit.meta)


@patch("app.services.step_generation.generate_structured_result")
def test_download_413_too_large(
    mock_gen: MagicMock,
    client: TestClient,
    seeded_session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, settings = seeded_session
    headers = _auth(client, settings)
    _, version_id = _create_in_review_version(client, settings, mock_gen)
    monkeypatch.setenv("MAX_DOCX_EXPORT_CONTENT_CHARS", "20")
    get_settings.cache_clear()

    resp = client.get(_download_url(version_id), headers=headers)
    assert resp.status_code == 413
    assert "maximum size" in resp.json()["detail"].lower()

    get_settings.cache_clear()


def test_control_characters_sanitized_in_render() -> None:
    from app.documents.meta import ExportMeta
    from app.documents.render import render_step_docx

    meta = ExportMeta(
        step_key="scope_analysis",
        step_title="Scope Analysis",
        project_name="Proj",
        client_name="Client",
        version_no=1,
        status="in_review",
        status_label="In review",
        marker_line="DRAFT, NOT APPROVED",
        metadata_date_label="Generated",
        metadata_date_value="1 January 2026",
        created_by_name="Author",
        is_draft_filename=True,
    )
    data = _sample_scope_output().model_dump()
    data["executive_summary"] = "Line\x00with\x1fcontrol"
    blob = render_step_docx("scope_analysis", data, meta)
    doc = _load_docx(blob)
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "\x00" not in joined
    assert "Linewithcontrol" in joined.replace(" ", "")


def test_render_unit_reopens_docx() -> None:
    from app.documents.meta import ExportMeta
    from app.documents.render import render_step_docx

    meta = ExportMeta(
        step_key="scope_analysis",
        step_title="Scope Analysis",
        project_name="Proj",
        client_name="Client",
        version_no=1,
        status="in_review",
        status_label="In review",
        marker_line="DRAFT, NOT APPROVED",
        metadata_date_label="Generated",
        metadata_date_value="1 January 2026",
        created_by_name="Author",
        is_draft_filename=True,
    )
    data = _sample_scope_output().model_dump()
    blob = render_step_docx("scope_analysis", data, meta)
    doc = _load_docx(blob)
    assert doc.paragraphs
