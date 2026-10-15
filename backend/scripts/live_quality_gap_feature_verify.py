"""Live dummy RFQ (30+ requirements) through Scope, Gap, Feature List. Fictional data only."""

from __future__ import annotations

import time
import uuid
from datetime import UTC, datetime
from io import BytesIO

import httpx
from docx import Document

BASE = "http://127.0.0.1:8080"

DUMMY_RFQ = """
Fictional training-only request from Contoso Example Ltd. Synthetic test data. Do not treat as real.

REQ-01 Must: Staff sign in to a customer portal with a hosted identity provider.
REQ-02 Must: Customers view a list of their orders with status and dates.
REQ-03 Must: Customers download PDF invoices for completed orders.
REQ-04 Must: Integrate order and invoice data from ExampleERP.
REQ-05 Must: Admins invite staff and assign roles.
REQ-06 Must: Customers reset a forgotten password.
REQ-07 Must: Show order status history on the order detail page.
REQ-08 Must: Block access to another customer's orders.
REQ-09 Should: Search orders by number, date range, and status.
REQ-10 Should: Email a notification when an order ships.
REQ-11 Should: Export orders to CSV for admins.
REQ-12 Should: Display stock availability copied from ExampleERP.
REQ-13 Should: Allow customers to update their billing address.
REQ-14 Should: Multi-location catalogue so each depot sees its own stock.
REQ-15 Could: Customers save favourite products.
REQ-16 Could: Dark mode for the portal.
REQ-17 Could: Optional SMS when an invoice is ready.
REQ-18 Could: Seasonal promotional banner managed by marketing.
REQ-19 Won't: Native mobile applications in this phase.
REQ-20 Won't: In-app live chat with agents.
REQ-21 Must: Record who changed an order status and when.
REQ-22 Must: Payments for outstanding invoices via a hosted checkout. Who owns the merchant account is not stated.
REQ-23 Must: Portal stays usable in a warehouse with intermittent connectivity.
REQ-24 Should: Webhook to ExampleERP when a customer confirms an address change.
REQ-25 Must: Named order states include Draft, Released, and Held with no definitions.
REQ-26 Should: Nightly import of open orders; timing window is not stated.
REQ-27 Must: Per-role screens for clerk, supervisor, and finance.
REQ-28 Should: Localisation hooks for a second language later.
REQ-29 Could: Simple sales report by depot and week.
REQ-30 Must: Empty and error states for the order list.
REQ-31 Should: Notification templates the client can edit.
REQ-32 Must: Scale target and which channels (web, kiosk, partner API) apply is not stated.
""".strip()


def login(client: httpx.Client, email: str, password: str) -> str:
    r = client.post("/auth/login", json={"email": email, "password": password})
    r.raise_for_status()
    return r.json()["access_token"]


def poll_version(client: httpx.Client, headers: dict[str, str], version_id: str) -> dict:
    status = "queued"
    detail: dict = {}
    for _ in range(180):
        resp = client.get(f"/step-versions/{version_id}", headers=headers)
        resp.raise_for_status()
        detail = resp.json()
        status = detail["status"]
        if status in {"in_review", "failed", "approved"}:
            return detail
        time.sleep(2)
    raise RuntimeError(f"timed out waiting for {version_id}, last status={status}")


def main() -> int:
    from app.config import get_settings

    settings = get_settings()
    email = settings.ADMIN_EMAIL
    password = settings.ADMIN_PASSWORD

    with httpx.Client(base_url=BASE, timeout=300.0) as client:
        token = login(client, email, password)
        headers = {"Authorization": f"Bearer {token}"}
        data = {
            "name": f"Fictional Quality RFQ {uuid.uuid4().hex[:8]}",
            "client_name": "Contoso Example Ltd (fictional)",
            "kind": "email",
            "received_at": datetime.now(UTC).isoformat(),
            "received_from": "fictional.client@example.test",
            "subject": "Fictional 32-requirement portal RFQ",
            "body": DUMMY_RFQ,
        }
        r = client.post("/projects", data=data, headers=headers)
        print("create_project", r.status_code)
        r.raise_for_status()
        pid = r.json()["project"]["id"]

        s1 = client.post(f"/projects/{pid}/steps/scope_analysis/generate", headers=headers)
        print("generate_scope", s1.status_code)
        s1.raise_for_status()
        scope = poll_version(client, headers, s1.json()["version"]["id"])
        print("scope", scope["status"], "prompt", scope.get("prompt_version"))
        if scope["status"] != "in_review":
            print("FAILED scope", scope.get("error"))
            return 1
        client.post(
            f"/step-versions/{scope['id']}/approve", headers=headers, json={}
        ).raise_for_status()

        s2 = client.post(f"/projects/{pid}/steps/gap_analysis/generate", headers=headers)
        print("generate_gap", s2.status_code)
        s2.raise_for_status()
        gap = poll_version(client, headers, s2.json()["version"]["id"])
        print("gap", gap["status"], "prompt", gap.get("prompt_version"))
        if gap["status"] != "in_review":
            print("FAILED gap", gap.get("error"))
            return 1
        gaps = (gap.get("content") or {}).get("gaps") or []
        without_evidence = sum(1 for item in gaps if not str(item.get("evidence") or "").strip())
        print("gap_count", len(gaps), "without_evidence", without_evidence)
        print("gap_notes", (gap.get("inputs") or {}).get("validation_notes"))
        client.post(
            f"/step-versions/{gap['id']}/approve", headers=headers, json={}
        ).raise_for_status()

        s3 = client.post(
            f"/projects/{pid}/steps/feature_list/generate",
            headers=headers,
            json={"inputs": {"team_notes": "Fictional: ExampleERP v4; hosted checkout is in scope."}},
        )
        print("generate_features", s3.status_code)
        s3.raise_for_status()
        features_version = poll_version(client, headers, s3.json()["version"]["id"])
        print(
            "feature",
            features_version["status"],
            "prompt",
            features_version.get("prompt_version"),
        )
        if features_version["status"] != "in_review":
            print("FAILED feature", features_version.get("error"))
            return 1
        content = features_version.get("content") or {}
        features = content.get("features") or []
        stats = content.get("stats") or {}
        print("feature_count", len(features))
        print("source_requirements_total", content.get("source_requirements_total"))
        print("stats", stats)
        print("warnings", content.get("warnings"))
        print("validation_notes", (features_version.get("inputs") or {}).get("validation_notes"))
        origins = stats.get("by_origin") or {}
        priorities = stats.get("by_priority") or {}
        print("implied", origins.get("implied"), "suggested", origins.get("suggested"))
        print("could", priorities.get("could"))
        dl = client.get(
            f"/step-versions/{features_version['id']}/download?format=docx",
            headers=headers,
        )
        print("docx", dl.status_code, len(dl.content))
        dl.raise_for_status()
        doc = Document(BytesIO(dl.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        header = doc.sections[0].header.paragraphs[0].text if doc.sections else ""
        print("docx_has_F-001", "F-001" in text)
        print("docx_self_approved", "Self-approved" in header or "Self-approved" in text)
        if len(gaps) < 10:
            print("WARN gap_count_below_10")
        if (priorities.get("could") or 0) == 0:
            print("WARN could_is_zero")
        if (origins.get("implied") or 0) + (origins.get("suggested") or 0) == 0:
            print("WARN no_implied_or_suggested")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print("FAILED", type(exc).__name__, exc)
        raise SystemExit(1) from exc
