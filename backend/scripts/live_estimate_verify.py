"""Live smoke: dummy project through Estimate. Use only fictional data."""

from __future__ import annotations

import csv
import subprocess
import tempfile
import time
import uuid
from datetime import UTC, datetime
from io import BytesIO, StringIO
from pathlib import Path

import httpx
from docx import Document
from openpyxl import load_workbook

BASE = "http://127.0.0.1:8080"

DUMMY_RFQ = """
Fictional training-only request from Contoso Example Ltd. Synthetic test data. Do not treat as real.

REQ-01 Must: Staff sign in to a customer portal with a hosted identity provider.
REQ-02 Must: Customers view a list of their orders with status and dates.
REQ-03 Must: Customers download PDF invoices for completed orders.
REQ-04 Must: Integrate order and invoice data from ExampleERP.
REQ-05 Must: Admins invite staff and assign roles.
REQ-06 Must: Customers reset a forgotten password.
REQ-07 Must: Show order status history on the order detail page.
REQ-08 Must: Block access to another customer's orders.
REQ-09 Should: Search orders by number, date range, and status.
REQ-10 Should: Email a notification when an order ships.
REQ-11 Should: Export orders to CSV for admins.
REQ-12 Should: Display stock availability copied from ExampleERP.
REQ-13 Should: Allow customers to update their billing address.
REQ-14 Should: Multi-location catalogue so each depot sees its own stock.
REQ-15 Could: Customers save favourite products.
REQ-16 Could: Dark mode for the portal.
REQ-17 Could: Optional SMS when an invoice is ready.
REQ-18 Could: Seasonal promotional banner managed by marketing.
REQ-19 Won't: Native mobile applications in this phase.
REQ-20 Won't: In-app live chat with agents.
REQ-21 Must: Record who changed an order status and when.
REQ-22 Must: Payments for outstanding invoices via a hosted checkout. Who owns the merchant account is not stated.
REQ-23 Must: Portal stays usable in a warehouse with intermittent connectivity.
REQ-24 Should: Webhook to ExampleERP when a customer confirms an address change.
REQ-25 Must: Named order states include Draft, Released, and Held with no definitions.
REQ-26 Should: Nightly import of open orders; timing window is not stated.
REQ-27 Must: Per-role screens for clerk, supervisor, and finance.
REQ-28 Should: Localisation hooks for a second language later.
REQ-29 Could: Simple sales report by depot and week.
REQ-30 Must: Empty and error states for the order list.
REQ-31 Should: Notification templates the client can edit.
REQ-32 Must: Scale target and which channels (web, kiosk, partner API) apply is not stated.
""".strip()

TYPICAL_TEAM = [
    {
        "role": "Project Manager",
        "discipline": "management",
        "mode": "duration_based",
        "headcount": 1,
        "hours_per_week": 40,
        "utilisation_pct": 100,
        "ai_gain_pct": 0,
    },
    {
        "role": "Developer",
        "discipline": "development",
        "mode": "estimated",
        "headcount": 4,
        "hours_per_week": 40,
        "utilisation_pct": 100,
        "ai_gain_pct": 50,
    },
    {
        "role": "QA Engineer",
        "discipline": "qa",
        "mode": "estimated",
        "headcount": 1,
        "hours_per_week": 40,
        "utilisation_pct": 100,
        "ai_gain_pct": 0,
    },
    {
        "role": "UI/UX Designer",
        "discipline": "design",
        "mode": "estimated",
        "headcount": 1,
        "hours_per_week": 40,
        "utilisation_pct": 100,
        "ai_gain_pct": 0,
    },
    {
        "role": "DevOps Engineer",
        "discipline": "devops",
        "mode": "percent_of_delivery",
        "headcount": 1,
        "hours_per_week": 40,
        "utilisation_pct": 100,
        "ai_gain_pct": 0,
        "overhead_pct": 5,
    },
]


def login(client: httpx.Client, email: str, password: str) -> str:
    r = client.post("/auth/login", json={"email": email, "password": password})
    r.raise_for_status()
    return r.json()["access_token"]


def poll_version(client: httpx.Client, headers: dict[str, str], version_id: str) -> dict:
    status = "queued"
    detail: dict = {}
    for _ in range(480):
        resp = client.get(f"/step-versions/{version_id}", headers=headers)
        resp.raise_for_status()
        detail = resp.json()
        status = detail["status"]
        if status in {"in_review", "failed", "approved"}:
            return detail
        time.sleep(2)
    raise RuntimeError(f"timed out waiting for {version_id}, last status={status}")


def generate_and_approve(
    client: httpx.Client,
    headers: dict[str, str],
    pid: str,
    step_key: str,
    inputs: dict | None = None,
) -> dict:
    payload: dict = {}
    if inputs is not None:
        payload["inputs"] = inputs
    resp = client.post(
        f"/projects/{pid}/steps/{step_key}/generate",
        headers=headers,
        json=payload or None,
    )
    print(f"generate_{step_key}", resp.status_code)
    resp.raise_for_status()
    detail = poll_version(client, headers, resp.json()["version"]["id"])
    print(f"{step_key}_status", detail["status"])
    if detail["status"] != "in_review":
        raise RuntimeError(f"{step_key} ended {detail['status']}: {detail.get('error')}")
    client.post(
        f"/step-versions/{detail['id']}/approve", headers=headers, json={}
    ).raise_for_status()
    return detail


def print_computed(label: str, content: dict) -> dict:
    computed = content.get("computed") if isinstance(content.get("computed"), dict) else {}
    totals = computed.get("totals") if isinstance(computed.get("totals"), dict) else {}
    phase_rows = computed.get("phase_rows") if isinstance(computed.get("phase_rows"), list) else []
    print(label, "hours_with_ai", totals.get("hours_with_ai"))
    print(label, "hours_without_ai", totals.get("hours_without_ai"))
    print(label, "recommended", totals.get("recommended_committed"))
    print(label, "weeks", totals.get("weeks"), "months", totals.get("months"))
    print(label, "ai_cap_factor", totals.get("project_ai_cap_factor"))
    print(label, "items", len(computed.get("item_export_rows") or []))
    print(label, "features", len(computed.get("feature_rows") or []))
    print(label, "work_packages", len(computed.get("work_package_rows") or []))
    print(label, "warnings", computed.get("warnings"))
    print(label, "cards", [(c.get("key"), c.get("value")) for c in (computed.get("cards") or []) if isinstance(c, dict)])
    for row in phase_rows:
        if not isinstance(row, dict):
            continue
        print(
            label,
            "phase",
            row.get("phase"),
            "hours_by_role",
            row.get("hours_by_role"),
            "weeks_planned",
            row.get("weeks_planned"),
        )
    assumptions = content.get("assumptions_and_solutions") or []
    print(label, "assumption_count", len(assumptions) if isinstance(assumptions, list) else 0)
    if isinstance(assumptions, list):
        for row in assumptions:
            if isinstance(row, dict):
                print(
                    "assumption",
                    row.get("gap_ref"),
                    row.get("topic"),
                    row.get("effort_impact"),
                )
    return totals


def recalc_xlsx(xlsx_bytes: bytes, dest: Path) -> Path:
    src = dest / "estimate-in.xlsx"
    src.write_bytes(xlsx_bytes)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "ods", "--outdir", str(dest), str(src)],
        check=True,
        capture_output=True,
        timeout=120,
    )
    ods = dest / "estimate-in.ods"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(dest), str(ods)],
        check=True,
        capture_output=True,
        timeout=120,
    )
    out = dest / "estimate-in.xlsx"
    # Conversion from ODS overwrites or creates a sibling; pick the newest xlsx besides src if needed.
    xlsx_files = sorted(dest.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return xlsx_files[0] if xlsx_files else out


def main() -> int:
    from app.config import get_settings

    settings = get_settings()
    email = settings.ADMIN_EMAIL
    password = settings.ADMIN_PASSWORD

    with httpx.Client(base_url=BASE, timeout=180.0) as client:
        token = login(client, email, password)
        headers = {"Authorization": f"Bearer {token}"}
        data = {
            "name": f"Fictional Estimate Pilot {uuid.uuid4().hex[:8]}",
            "client_name": "Contoso Example Ltd (fictional)",
            "kind": "email",
            "received_at": datetime.now(UTC).isoformat(),
            "received_from": "fictional.client@example.test",
            "subject": "Fictional 32-requirement portal RFQ",
            "body": DUMMY_RFQ,
        }
        r = client.post("/projects", data=data, headers=headers)
        print("create_project", r.status_code)
        r.raise_for_status()
        pid = r.json()["project"]["id"]

        steps = client.get(f"/projects/{pid}/steps", headers=headers).json()["steps"]
        est = next(s for s in steps if s["key"] == "estimate")
        print("step4_before", est["status"], est["blocked_reason"], est["export_formats"])

        generate_and_approve(client, headers, pid, "scope_analysis")
        generate_and_approve(client, headers, pid, "gap_analysis")
        generate_and_approve(
            client,
            headers,
            pid,
            "feature_list",
            {"team_notes": "Fictional client confirmed ExampleERP v4 and CSV export for admins."},
        )

        steps = client.get(f"/projects/{pid}/steps", headers=headers).json()["steps"]
        est = next(s for s in steps if s["key"] == "estimate")
        print("step4_ready", est["status"], est["can_generate"])

        team_v1 = [dict(row) for row in TYPICAL_TEAM]
        inputs_v1 = {
            "team": team_v1,
            "phases": ["mvp", "phase_1", "phase_2", "future_scope"],
            "domain_experience": "some_experience",
            "currency": "INR",
            "notes": "Fictional note: SSO can be Azure AD for this dummy portal.",
        }
        v1: dict | None = None
        for attempt in range(1, 5):
            s4 = client.post(
                f"/projects/{pid}/steps/estimate/generate",
                headers=headers,
                json={"inputs": inputs_v1},
            )
            print("generate_estimate", attempt, s4.status_code)
            if s4.status_code != 202:
                time.sleep(15)
                continue
            v1 = poll_version(client, headers, s4.json()["version"]["id"])
            print("estimate_v1", v1["status"], "error", v1.get("error"))
            if v1["status"] == "in_review":
                break
            time.sleep(15)
        if v1 is None or v1["status"] != "in_review":
            raise RuntimeError("estimate v1 did not reach in_review")
        content_v1 = v1.get("content") or {}
        totals_v1 = print_computed("v1", content_v1)
        print("validation_notes", (v1.get("inputs") or {}).get("validation_notes"))

        team_x2 = [dict(row) for row in TYPICAL_TEAM]
        team_x2[1]["headcount"] = 8
        rev = client.post(
            f"/step-versions/{v1['id']}/request-changes",
            headers=headers,
            json={
                "instructions": "Recalculate with the updated team.",
                "inputs": {**inputs_v1, "team": team_x2},
            },
        )
        print("request_changes_headcount", rev.status_code)
        rev.raise_for_status()
        v2 = poll_version(client, headers, rev.json()["version"]["id"])
        print("estimate_headcount", v2["status"], "error", v2.get("error"), "recompute", (v2.get("inputs") or {}).get("team_recompute"))
        if v2["status"] != "in_review":
            raise RuntimeError("estimate headcount recompute did not reach in_review")
        content_v2 = v2.get("content") or {}
        totals_v2 = print_computed("headcount_x2", content_v2)

        team_ai0 = [dict(row) for row in team_x2]
        team_ai0[1]["ai_gain_pct"] = 0
        rev0 = client.post(
            f"/step-versions/{v2['id']}/request-changes",
            headers=headers,
            json={
                "instructions": "Recalculate with the updated team.",
                "inputs": {**inputs_v1, "team": team_ai0},
            },
        )
        print("request_changes_ai0", rev0.status_code)
        rev0.raise_for_status()
        v_ai0 = poll_version(client, headers, rev0.json()["version"]["id"])
        print("estimate_ai0", v_ai0["status"], "error", v_ai0.get("error"))
        if v_ai0["status"] != "in_review":
            raise RuntimeError("estimate AI 0 recompute did not reach in_review")
        content_ai0 = v_ai0.get("content") or {}
        totals_ai0 = print_computed("ai0", content_ai0)
        print("schema_version", content_v2.get("schema_version"))
        print(
            "hours_without_ai_stable",
            totals_v1.get("hours_without_ai"),
            totals_v2.get("hours_without_ai"),
            totals_ai0.get("hours_without_ai"),
        )
        print(
            "headcount_v1_v2",
            (v1.get("inputs") or {}).get("user_inputs", {}).get("team", [{}])[1].get("headcount"),
            (v2.get("inputs") or {}).get("user_inputs", {}).get("team", [{}])[1].get("headcount"),
        )
        print("weeks_v1_v2", totals_v1.get("weeks"), totals_v2.get("weeks"))
        print("hours_ai_v1_v2", totals_v1.get("hours_with_ai"), totals_v2.get("hours_with_ai"))

        client.post(
            f"/step-versions/{v_ai0['id']}/approve", headers=headers, json={}
        ).raise_for_status()

        vid = v_ai0["id"]
        content_v2 = content_ai0
        totals_v2 = totals_ai0
        docx_resp = client.get(f"/step-versions/{vid}/download?format=docx", headers=headers)
        print("docx", docx_resp.status_code, len(docx_resp.content))
        docx_resp.raise_for_status()
        doc = Document(BytesIO(docx_resp.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        print("docx_has_health", "Estimate health" in text)
        print("docx_has_hours", "Recommended" in text or "hours with AI" in text.lower())
        print("docx_has_method", "without AI" in text.lower())

        xlsx_resp = client.get(f"/step-versions/{vid}/download?format=xlsx", headers=headers)
        print("xlsx", xlsx_resp.status_code, len(xlsx_resp.content))
        xlsx_resp.raise_for_status()
        wb = load_workbook(BytesIO(xlsx_resp.content))
        print("xlsx_sheets", wb.sheetnames)
        estimate_ws = wb["Estimate"]
        formulas = [
            cell.value
            for row in estimate_ws.iter_rows(min_row=2)
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        print("xlsx_formula_count", len(formulas))
        print("xlsx_sample_formula", formulas[0] if formulas else None)

        csv_resp = client.get(f"/step-versions/{vid}/download?format=csv", headers=headers)
        print("csv", csv_resp.status_code, len(csv_resp.content))
        csv_resp.raise_for_status()
        raw = csv_resp.content
        print("csv_bom", raw.startswith(b"\xef\xbb\xbf"))
        reader = csv.reader(StringIO(raw.decode("utf-8-sig")))
        rows = list(reader)
        print("csv_header", rows[0][:8] if rows else None)
        print("csv_first_rows", [row[:6] for row in rows[1:4]])

        with tempfile.TemporaryDirectory() as tmp:
            recalc_path = recalc_xlsx(xlsx_resp.content, Path(tmp))
            recalc = load_workbook(recalc_path, data_only=True)
            print("recalc_sheets", recalc.sheetnames)
            summary = recalc["Summary"]
            print("recalc_summary_b_formula_row", summary["B" + str(6 + len((content_v2.get("computed") or {}).get("cards") or []) + 1)].value)
            phases = recalc["Phases"]
            planned = list(phases.iter_rows(min_row=2, max_col=6, values_only=True))
            print("recalc_phase_rows", planned)
            computed_weeks = totals_v2.get("weeks")
            formula_weeks = summary[
                "B" + str(6 + len((content_v2.get("computed") or {}).get("cards") or []) + 1)
            ].value
            print("computed_weeks", computed_weeks, "recalc_committed_weeks", formula_weeks)

        print("ok", pid)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print("FAILED", type(exc).__name__, exc)
        raise SystemExit(1) from exc
