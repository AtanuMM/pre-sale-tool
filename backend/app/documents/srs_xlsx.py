"""SRS requirements register XLSX export."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.documents.meta import ExportMeta
from app.steps.schemas.srs import SrsStored


def build_srs_xlsx(content: dict[str, Any], meta: ExportMeta) -> bytes:
    stored = SrsStored.model_validate(content)
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    headers = ["ID", "Type", "Module", "Statement", "Priority", "Phase", "Basis", "Verification"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for fr in stored.functional_requirements:
        ws.append(
            [
                str(fr.id),
                "FR",
                str(fr.module),
                str(fr.statement),
                str(fr.priority),
                str(fr.phase),
                str(fr.basis),
                str(fr.verification_method),
            ]
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
