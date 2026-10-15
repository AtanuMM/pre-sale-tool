"""Estimate XLSX export with live formulas (openpyxl)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet

from app.config import get_settings
from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text


def _text_cell(ws: Worksheet, row: int, col: int, value: Any) -> None:
    text = "" if value is None else str(value)
    cell = ws.cell(row=row, column=col, value=sanitize_docx_text(text))
    cell.data_type = "s"
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _number_cell(ws: Worksheet, row: int, col: int, value: float | None) -> None:
    cell = ws.cell(row=row, column=col, value=0 if value is None else value)
    cell.number_format = "0.00"
    cell.alignment = Alignment(vertical="top")


def _formula_cell(ws: Worksheet, row: int, col: int, formula: str) -> None:
    value = formula if formula.startswith("=") else f"={formula}"
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = "0.00"
    cell.alignment = Alignment(vertical="top")


def _setup_sheet(ws: Worksheet, meta: ExportMeta) -> None:
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)
    ws.oddHeader.left.text = sanitize_docx_text(meta.marker_line)
    ws.oddHeader.right.text = sanitize_docx_text(f"{meta.project_name} v{meta.version_no}")
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _autosize(ws: Worksheet, cols: int, min_width: int = 12, max_width: int = 36) -> None:
    for idx in range(1, cols + 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = min(max_width, max(min_width, 14))


def _quote(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def build_estimate_xlsx(content: dict[str, Any], meta: ExportMeta) -> bytes:
    if isinstance(content, dict) and content.get("schema_version") == 2:
        return _build_estimate_xlsx_v2(content, meta)
    computed = content.get("computed") if isinstance(content.get("computed"), dict) else {}
    settings = get_settings()
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    team_ws = wb.create_sheet("Team")
    settings_ws = wb.create_sheet("Settings")
    estimate_ws = wb.create_sheet("Estimate")
    phases_ws = wb.create_sheet("Phases")
    components_ws = wb.create_sheet("Components")
    assumptions_ws = wb.create_sheet("Assumptions")
    exclusions_ws = wb.create_sheet("Exclusions")

    for ws in (
        summary,
        team_ws,
        settings_ws,
        estimate_ws,
        phases_ws,
        components_ws,
        assumptions_ws,
        exclusions_ws,
    ):
        _setup_sheet(ws, meta)

    # Settings named cells
    settings_ws["A1"] = "LeverageHigh"
    settings_ws["B1"] = settings.ESTIMATE_LEVERAGE_HIGH
    settings_ws["A2"] = "LeverageMedium"
    settings_ws["B2"] = settings.ESTIMATE_LEVERAGE_MEDIUM
    settings_ws["A3"] = "LeverageLow"
    settings_ws["B3"] = settings.ESTIMATE_LEVERAGE_LOW
    settings_ws["A4"] = "RiskHigh"
    settings_ws["B4"] = settings.ESTIMATE_RISK_HIGH
    settings_ws["A5"] = "RiskMedium"
    settings_ws["B5"] = settings.ESTIMATE_RISK_MEDIUM
    settings_ws["A6"] = "RiskLow"
    settings_ws["B6"] = settings.ESTIMATE_RISK_LOW
    settings_ws["A7"] = "Contingency"
    contig = computed.get("contingency_pct")
    settings_ws["B7"] = (float(contig) / 100.0) if isinstance(contig, (int, float)) else 0.15
    for row in range(1, 8):
        settings_ws.cell(row=row, column=2).number_format = "0.00"
    names = {
        "LeverageHigh": "$B$1",
        "LeverageMedium": "$B$2",
        "LeverageLow": "$B$3",
        "RiskHigh": "$B$4",
        "RiskMedium": "$B$5",
        "RiskLow": "$B$6",
        "Contingency": "$B$7",
    }
    for name, cell in names.items():
        wb.defined_names.add(DefinedName(name, attr_text=f"Settings!{cell}"))
    _autosize(settings_ws, 2)

    # Team sheet
    team_headers = [
        "Role",
        "Mode",
        "Headcount",
        "HoursPerWeek",
        "Utilisation",
        "AiGainPct",
        "HourlyRate",
        "OverheadPct",
    ]
    for col, header in enumerate(team_headers, start=1):
        team_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    team_rows = computed.get("team") if isinstance(computed.get("team"), list) else []
    for idx, row in enumerate(team_rows, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(team_ws, idx, 1, row.get("role"))
        _text_cell(team_ws, idx, 2, row.get("mode"))
        _number_cell(team_ws, idx, 3, row.get("headcount"))
        _number_cell(team_ws, idx, 4, row.get("hours_per_week"))
        _number_cell(team_ws, idx, 5, row.get("utilisation_pct"))
        _number_cell(team_ws, idx, 6, row.get("ai_gain_pct"))
        rate = row.get("hourly_rate")
        if isinstance(rate, (int, float)):
            _number_cell(team_ws, idx, 7, rate)
        oh = row.get("overhead_pct")
        if isinstance(oh, (int, float)):
            _number_cell(team_ws, idx, 8, oh)
    team_last = max(2, 1 + len(team_rows))
    _autosize(team_ws, 8)

    estimated_roles = [
        str(r) for r in (computed.get("estimated_roles") or []) if isinstance(r, str)
    ]
    export_rows = (
        computed.get("item_export_rows")
        if isinstance(computed.get("item_export_rows"), list)
        else []
    )

    # Estimate sheet
    est_headers = [
        "Status",
        "Phase",
        "Component",
        "ID",
        "Title",
        "Kind",
        "Confidence",
        "AILeverage",
    ]
    for role in estimated_roles:
        est_headers.extend([f"Base_{role}", f"AI_{role}", f"Buf_{role}"])
    est_headers.extend(["TotalAI", "TotalBuf"])
    for col, header in enumerate(est_headers, start=1):
        estimate_ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    def _leverage_expr(leverage_col: str) -> str:
        return (
            f'IF({leverage_col}="high",LeverageHigh,IF({leverage_col}="medium",'
            f"LeverageMedium,LeverageLow))"
        )

    def _risk_expr(conf_col: str) -> str:
        return (
            f'IF({conf_col}="high",RiskHigh,IF({conf_col}="medium",RiskMedium,RiskLow))'
        )

    for ridx, item in enumerate(export_rows, start=2):
        if not isinstance(item, dict):
            continue
        _text_cell(estimate_ws, ridx, 1, item.get("status"))
        _text_cell(estimate_ws, ridx, 2, item.get("phase"))
        _text_cell(estimate_ws, ridx, 3, item.get("component"))
        _text_cell(estimate_ws, ridx, 4, item.get("id"))
        _text_cell(estimate_ws, ridx, 5, item.get("title"))
        _text_cell(estimate_ws, ridx, 6, item.get("kind"))
        _text_cell(estimate_ws, ridx, 7, item.get("confidence"))
        _text_cell(estimate_ws, ridx, 8, item.get("ai_leverage"))
        per_role = item.get("per_role") if isinstance(item.get("per_role"), dict) else {}
        ai_letters: list[str] = []
        buf_letters: list[str] = []
        col = 9
        conf_col = f"G{ridx}"
        lev_col = f"H{ridx}"
        for role in estimated_roles:
            hours = per_role.get(role) if isinstance(per_role.get(role), dict) else {}
            base = hours.get("base") if isinstance(hours, dict) else 0
            _number_cell(estimate_ws, ridx, col, float(base or 0))
            base_letter = get_column_letter(col)
            ai_letter = get_column_letter(col + 1)
            buf_letter = get_column_letter(col + 2)
            role_q = _quote(role)
            ai_formula = (
                f"IFERROR(ROUND({base_letter}{ridx}*(1-IFERROR(INDEX(Team!$F$2:$F${team_last},"
                f"MATCH({role_q},Team!$A$2:$A${team_last},0)),0)/100*"
                f"{_leverage_expr(lev_col)}),2),0)"
            )
            buf_formula = (
                f"IFERROR(ROUND({ai_letter}{ridx}*(1+{_risk_expr(conf_col)}+Contingency),2),0)"
            )
            _formula_cell(estimate_ws, ridx, col + 1, ai_formula)
            _formula_cell(estimate_ws, ridx, col + 2, buf_formula)
            ai_letters.append(f"{ai_letter}{ridx}")
            buf_letters.append(f"{buf_letter}{ridx}")
            col += 3
        if ai_letters:
            _formula_cell(estimate_ws, ridx, col, f"SUM({','.join(ai_letters)})")
            _formula_cell(estimate_ws, ridx, col + 1, f"SUM({','.join(buf_letters)})")
        else:
            _number_cell(estimate_ws, ridx, col, 0)
            _number_cell(estimate_ws, ridx, col + 1, 0)
    _autosize(estimate_ws, max(8, len(est_headers)), min_width=10)
    est_last = max(2, 1 + len(export_rows))

    # Phases sheet — SUMIFS on Estimate buffered columns, MAX of weeks
    phase_headers = ["Phase", "Goal", "Rationale", "BufferedHours", "WeeksExact", "WeeksPlanned"]
    for col, header in enumerate(phase_headers, start=1):
        phases_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    phase_rows = computed.get("phase_rows") if isinstance(computed.get("phase_rows"), list) else []
    # TotalBuf column letter
    total_buf_col = get_column_letter(len(est_headers))
    phase_col = "B"
    for ridx, row in enumerate(phase_rows, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(phases_ws, ridx, 1, row.get("phase"))
        _text_cell(phases_ws, ridx, 2, row.get("goal"))
        _text_cell(phases_ws, ridx, 3, row.get("scope_rationale"))
        phase_q = _quote(str(row.get("phase") or ""))
        _formula_cell(
            phases_ws,
            ridx,
            4,
            f"SUMIF(Estimate!${phase_col}$2:${phase_col}${est_last},{phase_q},"
            f"Estimate!${total_buf_col}$2:${total_buf_col}${est_last})",
        )
        # weeks exact = max over estimated roles of buffered/capacity
        week_parts: list[str] = []
        for role_i, role in enumerate(estimated_roles):
            buf_col = get_column_letter(9 + role_i * 3 + 2)
            role_q = _quote(role)
            cap = (
                f"IFERROR(INDEX(Team!$C$2:$C${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))"
                f"*INDEX(Team!$D$2:$D${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))"
                f"*INDEX(Team!$E$2:$E${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))/100,0)"
            )
            hours = (
                f"SUMIF(Estimate!${phase_col}$2:${phase_col}${est_last},{phase_q},"
                f"Estimate!${buf_col}$2:${buf_col}${est_last})"
            )
            week_parts.append(f"IF(({cap})=0,0,({hours})/({cap}))")
        if week_parts:
            _formula_cell(phases_ws, ridx, 5, f"MAX({','.join(week_parts)})")
            _formula_cell(
                phases_ws,
                ridx,
                6,
                f"IF(E{ridx}<=0,0,CEILING(E{ridx},0.5))",
            )
        else:
            _number_cell(phases_ws, ridx, 5, 0)
            _number_cell(phases_ws, ridx, 6, 0)
    _autosize(phases_ws, 6)

    # Components
    comp_headers = ["Component", "BufferedHours"]
    for col, header in enumerate(comp_headers, start=1):
        components_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    component_totals = (
        computed.get("component_totals")
        if isinstance(computed.get("component_totals"), list)
        else []
    )
    comp_col_letter = "C"
    for ridx, row in enumerate(component_totals, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(components_ws, ridx, 1, row.get("component"))
        name_q = _quote(str(row.get("component") or ""))
        _formula_cell(
            components_ws,
            ridx,
            2,
            f"SUMIF(Estimate!${comp_col_letter}$2:${comp_col_letter}${est_last},{name_q},"
            f"Estimate!${total_buf_col}$2:${total_buf_col}${est_last})",
        )
    _autosize(components_ws, 2)

    # Summary
    _text_cell(summary, 1, 1, meta.marker_line)
    summary["A1"].font = Font(bold=True)
    _text_cell(summary, 2, 1, meta.step_title)
    _text_cell(summary, 3, 1, meta.project_name)
    _text_cell(summary, 4, 1, meta.client_name)
    cards = computed.get("cards") if isinstance(computed.get("cards"), list) else []
    for idx, card in enumerate(cards, start=6):
        if not isinstance(card, dict):
            continue
        _text_cell(summary, idx, 1, card.get("label"))
        _text_cell(summary, idx, 2, card.get("value"))
    _text_cell(summary, 6 + len(cards) + 1, 1, "Committed weeks (from Phases, excl. future_scope)")
    if phase_rows:
        last_phase_row = 1 + len(phase_rows)
        # Sum planned weeks where phase <> future_scope
        _formula_cell(
            summary,
            6 + len(cards) + 1,
            2,
            f'SUMIF(Phases!A2:A{last_phase_row},"<>future_scope",Phases!F2:F{last_phase_row})',
        )
    method = computed.get("method_note")
    if isinstance(method, str) and method:
        _text_cell(summary, 6 + len(cards) + 3, 1, method)
        summary.merge_cells(start_row=6 + len(cards) + 3, start_column=1, end_row=6 + len(cards) + 6, end_column=4)
    _autosize(summary, 2, min_width=28, max_width=60)

    # Assumptions
    ass_headers = [
        "Gap ref",
        "Topic",
        "Solution",
        "Rationale",
        "Effort impact",
        "Confirm with client",
    ]
    for col, header in enumerate(ass_headers, start=1):
        assumptions_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    assumptions = content.get("assumptions_and_solutions")
    if isinstance(assumptions, list):
        for ridx, row in enumerate(assumptions, start=2):
            if not isinstance(row, dict):
                continue
            _text_cell(assumptions_ws, ridx, 1, row.get("gap_ref"))
            _text_cell(assumptions_ws, ridx, 2, row.get("topic"))
            _text_cell(assumptions_ws, ridx, 3, row.get("solution"))
            _text_cell(assumptions_ws, ridx, 4, row.get("rationale"))
            _text_cell(assumptions_ws, ridx, 5, row.get("effort_impact"))
            _text_cell(assumptions_ws, ridx, 6, row.get("confirm_with_client"))
    _autosize(assumptions_ws, 6, max_width=40)

    exclusions = content.get("exclusions")
    exclusions_ws.cell(row=1, column=1, value="Exclusion").font = Font(bold=True)
    if isinstance(exclusions, list):
        for ridx, item in enumerate(exclusions, start=2):
            _text_cell(exclusions_ws, ridx, 1, item)
    _autosize(exclusions_ws, 1, max_width=60)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_estimate_xlsx_v2(content: dict[str, Any], meta: ExportMeta) -> bytes:
    computed = content.get("computed") if isinstance(content.get("computed"), dict) else {}
    settings = get_settings()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    team_ws = wb.create_sheet("Team")
    settings_ws = wb.create_sheet("Settings")
    estimate_ws = wb.create_sheet("Estimate")
    phases_ws = wb.create_sheet("Phases")
    components_ws = wb.create_sheet("Components")
    health_ws = wb.create_sheet("Health")
    assumptions_ws = wb.create_sheet("Assumptions")
    exclusions_ws = wb.create_sheet("Exclusions")
    for ws in (
        summary,
        team_ws,
        settings_ws,
        estimate_ws,
        phases_ws,
        components_ws,
        health_ws,
        assumptions_ws,
        exclusions_ws,
    ):
        _setup_sheet(ws, meta)

    totals = computed.get("totals") if isinstance(computed.get("totals"), dict) else {}
    factor = totals.get("project_ai_cap_factor")
    settings_rows = [
        ("CapHigh", settings.ESTIMATE_AI_CAP_HIGH),
        ("CapMedium", settings.ESTIMATE_AI_CAP_MEDIUM),
        ("CapLow", settings.ESTIMATE_AI_CAP_LOW),
        ("CapNone", settings.ESTIMATE_AI_CAP_NONE),
        ("BufferHigh", settings.ESTIMATE_BUFFER_HIGH),
        ("BufferMedium", settings.ESTIMATE_BUFFER_MEDIUM),
        ("BufferLow", settings.ESTIMATE_BUFFER_LOW),
        ("ProjectCapFactor", float(factor) if isinstance(factor, (int, float)) else 1.0),
    ]
    for idx, (name, value) in enumerate(settings_rows, start=1):
        settings_ws.cell(row=idx, column=1, value=name)
        settings_ws.cell(row=idx, column=2, value=value).number_format = "0.00"
        wb.defined_names.add(DefinedName(name, attr_text=f"Settings!$B${idx}"))
    _autosize(settings_ws, 2)

    team_headers = [
        "Role",
        "Discipline",
        "Mode",
        "Headcount",
        "HoursPerWeek",
        "Utilisation",
        "AiGainPct",
        "HourlyRate",
        "OverheadPct",
    ]
    for col, header in enumerate(team_headers, start=1):
        team_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    team_rows = computed.get("team") if isinstance(computed.get("team"), list) else []
    for idx, row in enumerate(team_rows, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(team_ws, idx, 1, row.get("role"))
        _text_cell(team_ws, idx, 2, row.get("discipline"))
        _text_cell(team_ws, idx, 3, row.get("mode"))
        _number_cell(team_ws, idx, 4, row.get("headcount"))
        _number_cell(team_ws, idx, 5, row.get("hours_per_week"))
        _number_cell(team_ws, idx, 6, row.get("utilisation_pct"))
        _number_cell(team_ws, idx, 7, row.get("ai_gain_pct"))
        rate = row.get("hourly_rate")
        if isinstance(rate, (int, float)):
            _number_cell(team_ws, idx, 8, rate)
        oh = row.get("overhead_pct")
        if isinstance(oh, (int, float)):
            _number_cell(team_ws, idx, 9, oh)
    team_last = max(2, 1 + len(team_rows))
    _autosize(team_ws, 9)

    estimated_roles = [
        str(r) for r in (computed.get("estimated_roles") or []) if isinstance(r, str)
    ]
    export_rows = (
        computed.get("item_export_rows")
        if isinstance(computed.get("item_export_rows"), list)
        else []
    )
    est_headers = [
        "Status",
        "Phase",
        "Component",
        "ID",
        "Title",
        "Kind",
        "Confidence",
        "AILeverage",
        "Best",
        "Likely",
        "Worst",
        "AILikely",
        "Recommended",
    ]
    for col, header in enumerate(est_headers, start=1):
        estimate_ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    def _cap_expr(lev_col: str) -> str:
        return (
            f'IF({lev_col}="high",CapHigh,IF({lev_col}="medium",CapMedium,'
            f'IF({lev_col}="low",CapLow,CapNone)))'
        )

    def _buf_expr(conf_col: str) -> str:
        return (
            f'IF({conf_col}="high",BufferHigh,IF({conf_col}="medium",BufferMedium,BufferLow))'
        )

    for ridx, item in enumerate(export_rows, start=2):
        if not isinstance(item, dict):
            continue
        _text_cell(estimate_ws, ridx, 1, item.get("status"))
        _text_cell(estimate_ws, ridx, 2, item.get("phase"))
        _text_cell(estimate_ws, ridx, 3, item.get("component"))
        _text_cell(estimate_ws, ridx, 4, item.get("id"))
        _text_cell(estimate_ws, ridx, 5, item.get("title"))
        _text_cell(estimate_ws, ridx, 6, item.get("kind"))
        _text_cell(estimate_ws, ridx, 7, item.get("confidence"))
        _text_cell(estimate_ws, ridx, 8, item.get("ai_leverage"))
        _number_cell(estimate_ws, ridx, 9, item.get("total_best"))
        _number_cell(estimate_ws, ridx, 10, item.get("total_likely"))
        _number_cell(estimate_ws, ridx, 11, item.get("total_worst"))
        # AILikely uses MIN of first development role's gain and leverage cap, times project cap.
        ai_formula = (
            f'IFERROR(ROUND(J{ridx}*(1-MIN(IFERROR(INDEX(Team!$G$2:$G${team_last},'
            f'MATCH("development",Team!$B$2:$B${team_last},0)),0)/100,'
            f"{_cap_expr('H' + str(ridx))})*ProjectCapFactor),2),J{ridx})"
        )
        rec_formula = f"IFERROR(ROUND(L{ridx}*(1+{_buf_expr('G' + str(ridx))}),2),0)"
        _formula_cell(estimate_ws, ridx, 12, ai_formula)
        _formula_cell(estimate_ws, ridx, 13, rec_formula)
    _autosize(estimate_ws, 13, min_width=10)
    est_last = max(2, 1 + len(export_rows))

    phase_headers = ["Phase", "Goal", "Rationale", "RecommendedHours", "WeeksExact", "WeeksPlanned"]
    for col, header in enumerate(phase_headers, start=1):
        phases_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    phase_rows = computed.get("phase_rows") if isinstance(computed.get("phase_rows"), list) else []
    for ridx, row in enumerate(phase_rows, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(phases_ws, ridx, 1, row.get("phase"))
        _text_cell(phases_ws, ridx, 2, row.get("goal"))
        _text_cell(phases_ws, ridx, 3, row.get("scope_rationale"))
        phase_q = _quote(str(row.get("phase") or ""))
        _formula_cell(
            phases_ws,
            ridx,
            4,
            f"SUMIF(Estimate!$B$2:$B${est_last},{phase_q},Estimate!$M$2:$M${est_last})",
        )
        week_parts: list[str] = []
        for role in estimated_roles:
            role_q = _quote(role)
            cap = (
                f"IFERROR(INDEX(Team!$D$2:$D${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))"
                f"*INDEX(Team!$E$2:$E${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))"
                f"*INDEX(Team!$F$2:$F${team_last},MATCH({role_q},Team!$A$2:$A${team_last},0))/100,0)"
            )
            week_parts.append(f"IF(({cap})=0,0,D{ridx}/({cap})/{max(len(estimated_roles), 1)})")
        if week_parts:
            _formula_cell(phases_ws, ridx, 5, f"MAX({','.join(week_parts)})")
            _formula_cell(phases_ws, ridx, 6, f"IF(E{ridx}<=0,0,CEILING(E{ridx},0.5))")
        else:
            _number_cell(phases_ws, ridx, 5, 0)
            _number_cell(phases_ws, ridx, 6, 0)
    _autosize(phases_ws, 6)

    components_ws.cell(row=1, column=1, value="Component").font = Font(bold=True)
    components_ws.cell(row=1, column=2, value="RecommendedHours").font = Font(bold=True)
    component_totals = (
        computed.get("component_totals")
        if isinstance(computed.get("component_totals"), list)
        else []
    )
    for ridx, row in enumerate(component_totals, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(components_ws, ridx, 1, row.get("component"))
        name_q = _quote(str(row.get("component") or ""))
        _formula_cell(
            components_ws,
            ridx,
            2,
            f"SUMIF(Estimate!$C$2:$C${est_last},{name_q},Estimate!$M$2:$M${est_last})",
        )
    _autosize(components_ws, 2)

    health_ws.cell(row=1, column=1, value="Check").font = Font(bold=True)
    health_ws.cell(row=1, column=2, value="Value").font = Font(bold=True)
    health_ws.cell(row=1, column=3, value="Threshold").font = Font(bold=True)
    health_ws.cell(row=1, column=4, value="Status").font = Font(bold=True)
    health_rows = computed.get("health_rows") if isinstance(computed.get("health_rows"), list) else []
    for ridx, row in enumerate(health_rows, start=2):
        if not isinstance(row, dict):
            continue
        _text_cell(health_ws, ridx, 1, row.get("check"))
        _text_cell(health_ws, ridx, 2, row.get("value"))
        _text_cell(health_ws, ridx, 3, row.get("threshold"))
        _text_cell(health_ws, ridx, 4, row.get("status"))
    _autosize(health_ws, 4, min_width=18, max_width=50)

    _text_cell(summary, 1, 1, meta.marker_line)
    summary["A1"].font = Font(bold=True)
    _text_cell(summary, 2, 1, meta.step_title)
    _text_cell(summary, 3, 1, meta.project_name)
    _text_cell(summary, 4, 1, meta.client_name)
    cards = computed.get("cards") if isinstance(computed.get("cards"), list) else []
    for idx, card in enumerate(cards, start=6):
        if not isinstance(card, dict):
            continue
        _text_cell(summary, idx, 1, card.get("label"))
        _text_cell(summary, idx, 2, card.get("value"))
    method = computed.get("method_note")
    if isinstance(method, str) and method:
        _text_cell(summary, 6 + len(cards) + 2, 1, method)
        summary.merge_cells(
            start_row=6 + len(cards) + 2,
            start_column=1,
            end_row=6 + len(cards) + 5,
            end_column=4,
        )
    _autosize(summary, 2, min_width=28, max_width=60)

    ass_headers = ["Gap ref", "Topic", "Solution", "Rationale", "Effort impact", "Basis"]
    for col, header in enumerate(ass_headers, start=1):
        assumptions_ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    assumptions = content.get("assumptions_and_solutions")
    if isinstance(assumptions, list):
        for ridx, row in enumerate(assumptions, start=2):
            if not isinstance(row, dict):
                continue
            _text_cell(assumptions_ws, ridx, 1, row.get("gap_ref"))
            _text_cell(assumptions_ws, ridx, 2, row.get("topic"))
            _text_cell(assumptions_ws, ridx, 3, row.get("solution"))
            _text_cell(assumptions_ws, ridx, 4, row.get("rationale"))
            _text_cell(assumptions_ws, ridx, 5, row.get("effort_impact"))
            _text_cell(assumptions_ws, ridx, 6, row.get("basis"))
    _autosize(assumptions_ws, 6, max_width=40)

    exclusions_ws.cell(row=1, column=1, value="Exclusion").font = Font(bold=True)
    exclusions = content.get("exclusions")
    if isinstance(exclusions, list):
        for ridx, item in enumerate(exclusions, start=2):
            _text_cell(exclusions_ws, ridx, 1, item)
    _autosize(exclusions_ws, 1, max_width=60)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

