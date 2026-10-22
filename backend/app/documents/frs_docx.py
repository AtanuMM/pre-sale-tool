"""FRS DOCX export."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.documents.docx_builder import _set_run_font, create_export_document
from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text
from app.steps.schemas.frs import FrsStored


def render_frs_docx(content: dict[str, Any], meta: ExportMeta) -> bytes:
    stored = FrsStored.model_validate(content)
    marker = "Approved for issue" if meta.status == "approved" else meta.marker_line
    doc = create_export_document(
        ExportMeta(
            step_key=meta.step_key,
            step_title=meta.step_title,
            project_name=meta.project_name,
            client_name=meta.client_name,
            version_no=meta.version_no,
            status=meta.status,
            status_label=meta.status_label,
            marker_line=marker,
            metadata_date_label=meta.metadata_date_label,
            metadata_date_value=meta.metadata_date_value,
            created_by_name=meta.created_by_name,
            is_draft_filename=meta.is_draft_filename,
        )
    )
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    if marker:
        run = header.add_run(sanitize_docx_text(marker))
        _set_run_font(run, size_pt=9, bold=True)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        meta.client_name,
        meta.project_name,
        "Functional Requirements Specification",
        stored.assembled.document_reference,
        meta.metadata_date_label + ": " + meta.metadata_date_value,
    ):
        run = cover.add_run(sanitize_docx_text(line) + "\n")
        _set_run_font(run, size_pt=11)

    doc.add_page_break()
    h = doc.add_paragraph("1 Introduction")
    _set_run_font(h.runs[0], size_pt=14, bold=True)
    p = doc.add_paragraph(
        sanitize_docx_text(
            "This document describes functional behaviour for the committed release scope."
        )
    )
    _set_run_font(p.runs[0], size_pt=11)

    doc.add_page_break()
    h2 = doc.add_paragraph("3 Functional specifications")
    _set_run_font(h2.runs[0], size_pt=14, bold=True)
    for spec in stored.specifications[:200]:
        doc.add_paragraph(
            sanitize_docx_text(f"{spec.id} — {spec.title} ({spec.module})"),
            style="Heading 2",
        )
        body = doc.add_paragraph(sanitize_docx_text(spec.summary))
        _set_run_font(body.runs[0], size_pt=11)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
