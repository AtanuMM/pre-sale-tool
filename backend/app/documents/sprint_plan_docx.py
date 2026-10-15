"""Sprint Plan DOCX export."""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.documents.docx_builder import _set_run_font, create_export_document
from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text
from app.steps.schemas.sprint_plan import SprintPlanStored


def render_sprint_plan_docx(content: dict[str, Any], meta: ExportMeta) -> bytes:
    stored = SprintPlanStored.model_validate(content)
    show_hours = stored.viewer.comparison.get("show_hours") is True
    doc = create_export_document(meta)
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        meta.client_name,
        meta.project_name,
        "Sprint Plan",
        stored.assembled.document_reference,
        meta.metadata_date_label + ": " + meta.metadata_date_value,
    ):
        run = cover.add_run(sanitize_docx_text(str(line)) + "\n")
        _set_run_font(run, size_pt=11)

    doc.add_page_break()
    h = doc.add_paragraph(style="Heading 1")
    run = h.add_run("1 Introduction")
    _set_run_font(run, size_pt=14, bold=True)
    if stored.passes.goals and stored.passes.goals.executive_summary:
        p = doc.add_paragraph()
        r = p.add_run(sanitize_docx_text(stored.passes.goals.executive_summary))
        _set_run_font(r, size_pt=11)

    h2 = doc.add_paragraph(style="Heading 1")
    run2 = h2.add_run("5 Sprint plan")
    _set_run_font(run2, size_pt=14, bold=True)
    for sprint in stored.assembled.schedule_sprints:
        title = doc.add_paragraph(style="Heading 2")
        tr = title.add_run(sanitize_docx_text(f"{sprint.get('name')} ({sprint.get('phase')})"))
        _set_run_font(tr, size_pt=12, bold=True)
        for story in stored.stories:
            if story.sprint_number != sprint.get("number"):
                continue
            row = doc.add_paragraph()
            text = f"{story.id}: {story.title}"
            if show_hours and story.hours_by_role:
                text += " — " + ", ".join(f"{k} {v:g}h" for k, v in story.hours_by_role.items())
            rr = row.add_run(sanitize_docx_text(text))
            _set_run_font(rr, size_pt=10)

    from io import BytesIO

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
