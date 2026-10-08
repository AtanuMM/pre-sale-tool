"""Word document shell: page setup, header/footer fields, core properties."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text


def _set_run_font(run, *, size_pt: float, bold: bool = False) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    run.bold = bold


def _add_field_run(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    run._r.append(instr_el)

    run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    run = paragraph.add_run("1")
    _set_run_font(run, size_pt=9)

    run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def _configure_page(section) -> None:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def _apply_default_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(14)
    h1.font.bold = True


def create_export_document(meta: ExportMeta) -> Document:
    doc = Document()
    _apply_default_style(doc)
    section = doc.sections[0]
    _configure_page(section)

    props = doc.core_properties
    props.title = sanitize_docx_text(f"{meta.step_title} – {meta.project_name}")
    props.author = "ScopeDesk"
    props.subject = sanitize_docx_text(meta.project_name)

    marker = sanitize_docx_text(meta.marker_line)
    header = section.header
    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.clear()
    if marker:
        run = header_para.add_run(marker)
        _set_run_font(run, size_pt=9, bold=True)

    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.clear()
    prefix = footer_para.add_run(
        sanitize_docx_text(f"ScopeDesk · {meta.project_name} · Page ")
    )
    _set_run_font(prefix, size_pt=9)
    _add_field_run(footer_para, "PAGE")
    mid = footer_para.add_run(" of ")
    _set_run_font(mid, size_pt=9)
    _add_field_run(footer_para, "NUMPAGES")

    if marker:
        marker_para = doc.add_paragraph()
        marker_run = marker_para.add_run(marker)
        _set_run_font(marker_run, size_pt=12, bold=True)

    title = doc.add_paragraph(style="Heading 1")
    title_run = title.add_run(sanitize_docx_text(meta.step_title))
    _set_run_font(title_run, size_pt=16, bold=True)

    doc.add_paragraph(sanitize_docx_text(meta.project_name))
    doc.add_paragraph(sanitize_docx_text(meta.client_name))

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = (
        ("Version", str(meta.version_no)),
        ("Status", meta.status_label),
        (meta.metadata_date_label, meta.metadata_date_value),
        ("Created by", meta.created_by_name),
    )
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = sanitize_docx_text(label)
        table.rows[idx].cells[1].text = sanitize_docx_text(value)

    doc.add_paragraph()
    return doc
