"""Generic step content → DOCX bytes."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt
from fastapi import HTTPException, status

from app.config import get_settings
from app.documents.docx_builder import create_export_document
from app.documents.layouts import layout_for_step
from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text
from app.steps.layout_config import (
    LayoutColumn,
    SectionKind,
    SectionLayoutSpec,
    resolve_section_value,
)
from app.steps.refs import ref_label


def content_char_weight(value: Any) -> int:
    if isinstance(value, str):
        return len(value)
    if isinstance(value, dict):
        return sum(content_char_weight(v) for v in value.values())
    if isinstance(value, list):
        return sum(content_char_weight(item) for item in value)
    return 0


def assert_content_within_export_limit(content: dict[str, Any]) -> None:
    settings = get_settings()
    total = content_char_weight(content)
    if total > settings.MAX_DOCX_EXPORT_CONTENT_CHARS:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=(
                "Step version content exceeds the maximum size allowed for export "
                f"({settings.MAX_DOCX_EXPORT_CONTENT_CHARS} characters)"
            ),
        )


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return sanitize_docx_text(value)
    if isinstance(value, list):
        parts = [_cell_text(item) for item in value]
        return sanitize_docx_text(", ".join(p for p in parts if p))
    return sanitize_docx_text(str(value))


def _humanize_token(value: str) -> str:
    pretty = value.replace("_", " ")
    if not pretty:
        return pretty
    return pretty[0].upper() + pretty[1:]


def _humanize_field(name: str) -> str:
    trimmed = name.removesuffix("_ids").removesuffix("_id").removesuffix("_refs")
    return _humanize_token(trimmed.replace("_", " "))


def _add_section_heading(doc: Document, title: str, *, level: int = 1) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    para = doc.add_paragraph(style=style)
    run = para.add_run(sanitize_docx_text(title))
    run.font.name = "Calibri"
    run.font.size = Pt(14 if level == 1 else 12)
    run.bold = True


def _render_text(doc: Document, raw: Any) -> None:
    if not isinstance(raw, str) or not raw.strip():
        return
    doc.add_paragraph(sanitize_docx_text(raw))


def _render_list(doc: Document, raw: Any, layout: SectionLayoutSpec) -> None:
    if not isinstance(raw, list):
        return
    for index, item in enumerate(raw, start=1):
        if not isinstance(item, str) or not item.strip():
            continue
        text = sanitize_docx_text(item)
        if layout.ref_prefix:
            text = f"{ref_label(layout.ref_prefix, index)}  {text}"
        doc.add_paragraph(text, style="List Bullet")


def _render_numbered_list(doc: Document, raw: Any, layout: SectionLayoutSpec) -> None:
    if not isinstance(raw, list):
        return
    for index, item in enumerate(raw, start=1):
        if not isinstance(item, str) or not item.strip():
            continue
        text = sanitize_docx_text(item)
        if layout.ref_prefix:
            text = f"{ref_label(layout.ref_prefix, index)}  {text}"
        doc.add_paragraph(text, style="List Number")


def _render_table(doc: Document, raw: Any, layout: SectionLayoutSpec) -> None:
    if not isinstance(raw, list) or not layout.columns:
        return
    objects = [row for row in raw if isinstance(row, dict)]
    if not objects:
        return
    table = doc.add_table(rows=1 + len(objects), cols=len(layout.columns))
    table.style = "Table Grid"
    for col_idx, col in enumerate(layout.columns):
        table.rows[0].cells[col_idx].text = sanitize_docx_text(col.header)
    for row_idx, obj in enumerate(objects, start=1):
        for col_idx, col in enumerate(layout.columns):
            table.rows[row_idx].cells[col_idx].text = _cell_text(obj.get(col.key))


def _chip_values(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [_cell_text(item) for item in value if _cell_text(item)]
    text = _cell_text(value)
    return [text] if text else []


def _render_record_block(
    doc: Document, obj: dict[str, Any], layout: SectionLayoutSpec, *, index: int
) -> None:
    chips: list[str] = []
    if layout.ref_prefix:
        chips.append(ref_label(layout.ref_prefix, index))
    if layout.id_field:
        id_val = _cell_text(obj.get(layout.id_field))
        if id_val:
            chips.append(id_val)
    if chips:
        doc.add_paragraph(sanitize_docx_text(" ".join(chips)))
    badge_parts: list[str] = []
    for badge_key in layout.badge_fields:
        val = obj.get(badge_key)
        text = _cell_text(val)
        if not text:
            continue
        if isinstance(val, list):
            pretty = ", ".join(_humanize_token(part.strip()) for part in text.split(",") if part.strip())
            badge_parts.append(f"{_humanize_field(badge_key)}: {pretty}")
        else:
            badge_parts.append(f"{_humanize_field(badge_key)}: {_humanize_token(text)}")
    if badge_parts:
        doc.add_paragraph(sanitize_docx_text(" · ".join(badge_parts)))
    for field in layout.record_fields:
        val = obj.get(field.key)
        if val is None or (isinstance(val, str) and not val.strip()):
            continue
        line = f"{field.header}: {_cell_text(val)}"
        doc.add_paragraph(sanitize_docx_text(line))
    for chip_field in layout.chip_fields:
        values = _chip_values(obj.get(chip_field))
        if not values:
            continue
        line = f"{_humanize_field(chip_field)}: {', '.join(values)}"
        doc.add_paragraph(sanitize_docx_text(line))
    doc.add_paragraph()


def _grouped_records(
    objects: list[tuple[int, dict[str, Any]]],
    layout: SectionLayoutSpec,
    content: dict[str, Any],
) -> list[tuple[str, str, list[tuple[int, dict[str, Any]]]]]:
    if not layout.group_by or layout.group_source is None:
        return [("", "", objects)]
    source = content.get(layout.group_source.section)
    order: list[tuple[str, str]] = []
    seen: set[str] = set()
    if isinstance(source, list):
        for item in source:
            if not isinstance(item, dict):
                continue
            name = str(item.get(layout.group_source.key_field) or "").strip()
            if not name or name.lower() in seen:
                continue
            seen.add(name.lower())
            intro = str(item.get(layout.group_source.intro_field) or "").strip()
            order.append((name, intro))
    buckets: dict[str, list[tuple[int, dict[str, Any]]]] = {name: [] for name, _ in order}
    extras: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for indexed in objects:
        name = str(indexed[1].get(layout.group_by) or "").strip() or "Other"
        key = next((n for n, _ in order if n.lower() == name.lower()), None)
        if key is not None:
            buckets[key].append(indexed)
        else:
            extras.setdefault(name, []).append(indexed)
    groups: list[tuple[str, str, list[tuple[int, dict[str, Any]]]]] = []
    intro_by_name = {n: intro for n, intro in order}
    for name, _intro in order:
        items = buckets.get(name) or []
        if items:
            groups.append((name, intro_by_name.get(name, ""), items))
    for name, items in extras.items():
        groups.append((name, "", items))
    return groups or [("", "", objects)]


def _render_records(
    doc: Document, raw: Any, layout: SectionLayoutSpec, content: dict[str, Any]
) -> None:
    if not isinstance(raw, list) or not layout.record_fields:
        return
    objects = [
        (index, row)
        for index, row in enumerate(raw, start=1)
        if isinstance(row, dict)
    ]
    if not objects:
        return
    for heading, intro, items in _grouped_records(objects, layout, content):
        if heading:
            _add_section_heading(doc, heading, level=2)
        if intro:
            doc.add_paragraph(sanitize_docx_text(intro))
        for index, obj in items:
            _render_record_block(doc, obj, layout, index=index)


def _render_stats(doc: Document, raw: Any, layout: SectionLayoutSpec) -> None:
    if not isinstance(raw, dict):
        return
    parts: list[str] = []
    for group in layout.stats_groups:
        mapping = raw.get(group.key)
        if not isinstance(mapping, dict) or not mapping:
            continue
        inner = ", ".join(
            f"{_humanize_token(str(k))} {v}" for k, v in mapping.items()
        )
        parts.append(f"{group.header}: {inner}")
    if parts:
        doc.add_paragraph(sanitize_docx_text(" · ".join(parts)))


def _render_computed_cards(doc: Document, raw: Any) -> None:
    items = raw if isinstance(raw, list) else []
    for item in items:
        if not isinstance(item, dict):
            continue
        label = _cell_text(item.get("label"))
        value = _cell_text(item.get("value"))
        if not label and not value:
            continue
        doc.add_paragraph(sanitize_docx_text(f"{label}: {value}"))


def _render_warning_strip(doc: Document, raw: Any) -> bool:
    items = raw if isinstance(raw, list) else []
    lines: list[str] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        count = item.get("count")
        label = _cell_text(item.get("label"))
        if not isinstance(count, (int, float)) or count <= 0:
            continue
        lines.append(f"{label}: {int(count)}")
    if not lines:
        return False
    doc.add_paragraph(sanitize_docx_text(" · ".join(lines)))
    return True


def _render_computed_table(doc: Document, raw: Any, layout: SectionLayoutSpec) -> None:
    if not isinstance(raw, list):
        return
    objects = [row for row in raw if isinstance(row, dict)]
    if not objects:
        return
    columns = list(layout.columns)
    if not columns:
        keys = [k for k in objects[0] if k != "per_role"]
        columns = [LayoutColumn(str(k), _humanize_token(str(k))) for k in keys]
    if not columns:
        return
    table = doc.add_table(rows=1 + len(objects), cols=len(columns))
    table.style = "Table Grid"
    for col_idx, col in enumerate(columns):
        table.rows[0].cells[col_idx].text = sanitize_docx_text(col.header)
    for row_idx, obj in enumerate(objects, start=1):
        for col_idx, col in enumerate(columns):
            table.rows[row_idx].cells[col_idx].text = _cell_text(obj.get(col.key))


def _render_section(doc: Document, content: dict[str, Any], layout: SectionLayoutSpec) -> None:
    raw = resolve_section_value(content, layout)
    if raw is None:
        return
    if layout.kind == SectionKind.TEXT or layout.kind == SectionKind.METHOD_NOTE:
        if not isinstance(raw, str) or not raw.strip():
            return
    elif layout.kind == SectionKind.LIST or layout.kind == SectionKind.NUMBERED_LIST:
        if not isinstance(raw, list) or not any(isinstance(x, str) and x.strip() for x in raw):
            return
    elif layout.kind == SectionKind.TABLE and (
        not isinstance(raw, list) or not any(isinstance(x, dict) for x in raw)
    ) or layout.kind == SectionKind.RECORDS and (
        not isinstance(raw, list) or not any(isinstance(x, dict) for x in raw)
    ) or layout.kind == SectionKind.STATS and (not isinstance(raw, dict) or not raw) or layout.kind == SectionKind.COMPUTED_CARDS and (
        not isinstance(raw, list) or not raw
    ) or layout.kind == SectionKind.COMPUTED_TABLE and (
        not isinstance(raw, list) or not any(isinstance(x, dict) for x in raw)
    ):
        return
    elif layout.kind == SectionKind.WARNING_STRIP:
        items = raw if isinstance(raw, list) else []
        if not any(
            isinstance(x, dict) and isinstance(x.get("count"), (int, float)) and x.get("count") > 0
            for x in items
        ):
            return
    _add_section_heading(doc, layout.title)
    if layout.kind == SectionKind.TEXT or layout.kind == SectionKind.METHOD_NOTE:
        _render_text(doc, raw)
    elif layout.kind == SectionKind.LIST:
        _render_list(doc, raw, layout)
    elif layout.kind == SectionKind.NUMBERED_LIST:
        _render_numbered_list(doc, raw, layout)
    elif layout.kind == SectionKind.TABLE:
        _render_table(doc, raw, layout)
    elif layout.kind == SectionKind.RECORDS:
        _render_records(doc, raw, layout, content)
    elif layout.kind == SectionKind.STATS:
        _render_stats(doc, raw, layout)
    elif layout.kind == SectionKind.COMPUTED_CARDS:
        _render_computed_cards(doc, raw)
    elif layout.kind == SectionKind.COMPUTED_TABLE:
        _render_computed_table(doc, raw, layout)
    elif layout.kind == SectionKind.WARNING_STRIP:
        _render_warning_strip(doc, raw)
    doc.add_paragraph()


def render_step_docx(step_key: str, content: dict[str, Any], meta: ExportMeta) -> bytes:
    layout = layout_for_step(step_key, content)
    if layout is None:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="DOCX export is not available for this step",
        )
    assert_content_within_export_limit(content)
    doc = create_export_document(meta)
    for section in layout:
        _render_section(doc, content, section)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
