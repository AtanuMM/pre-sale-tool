"""Statement of Work DOCX export."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.documents.docx_builder import _set_run_font, create_export_document
from app.documents.meta import ExportMeta
from app.documents.sanitize import sanitize_docx_text
from app.steps.schemas.sow import SowStored


def _add_heading(doc, text: str, *, level: int = 1) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    para = doc.add_paragraph(style=style)
    run = para.add_run(sanitize_docx_text(text))
    _set_run_font(run, size_pt=14 if level == 1 else 12, bold=True)


def _add_rich_paragraph(doc, text: str) -> None:
    para = doc.add_paragraph()
    parts = sanitize_docx_text(text).split("[TBC:")
    if len(parts) == 1:
        run = para.add_run(parts[0])
        _set_run_font(run, size_pt=11)
        return
    run = para.add_run(parts[0])
    _set_run_font(run, size_pt=11)
    for chunk in parts[1:]:
        end = chunk.find("]")
        if end == -1:
            run = para.add_run("[TBC:" + chunk)
            _set_run_font(run, size_pt=11, bold=True)
            continue
        tbc = "[TBC:" + chunk[: end + 1]
        run = para.add_run(tbc)
        _set_run_font(run, size_pt=11, bold=True)
        rest = chunk[end + 1 :]
        if rest:
            run = para.add_run(rest)
            _set_run_font(run, size_pt=11)


def _add_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, title in enumerate(headers):
        hdr[idx].text = sanitize_docx_text(title)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = sanitize_docx_text(val)


def _sow_marker_line(meta: ExportMeta) -> str:
    if meta.status == "approved":
        return "Approved for issue"
    return meta.marker_line


def render_sow_docx(content: dict[str, Any], meta: ExportMeta) -> bytes:
    stored = SowStored.model_validate(content)
    doc = create_export_document(
        ExportMeta(
            step_key=meta.step_key,
            step_title=meta.step_title,
            project_name=meta.project_name,
            client_name=meta.client_name,
            version_no=meta.version_no,
            status=meta.status,
            status_label=meta.status_label,
            marker_line=_sow_marker_line(meta),
            metadata_date_label=meta.metadata_date_label,
            metadata_date_value=meta.metadata_date_value,
            created_by_name=meta.created_by_name,
            is_draft_filename=meta.is_draft_filename,
        )
    )
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    marker = _sow_marker_line(meta)
    if marker:
        run = header.add_run(sanitize_docx_text(marker))
        _set_run_font(run, size_pt=9, bold=True)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    inputs = stored.assembled
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        meta.client_name,
        meta.project_name,
        stored.narrative.document_title,
        inputs.document_reference,
        f"Version {meta.version_no}",
        inputs.issue_date,
        inputs.status_text,
    ):
        run = cover.add_run(sanitize_docx_text(str(line)) + "\n")
        _set_run_font(run, size_pt=12, bold=line == stored.narrative.document_title)

    doc.add_page_break()
    _add_heading(doc, "Contents", level=1)
    for item in inputs.contents_list:
        doc.add_paragraph(sanitize_docx_text(item), style="List Number")

    doc.add_page_break()
    _add_heading(doc, "1 Introduction", level=1)
    _add_rich_paragraph(doc, stored.narrative.purpose)
    _add_rich_paragraph(doc, stored.narrative.background)

    _add_heading(doc, "2 Objectives and success criteria", level=1)
    for obj in stored.narrative.objectives:
        doc.add_paragraph(sanitize_docx_text(obj), style="List Bullet")
    for crit in stored.narrative.success_criteria:
        doc.add_paragraph(sanitize_docx_text(crit), style="List Bullet")

    _add_heading(doc, "3 Scope of work", level=1)
    _add_rich_paragraph(doc, stored.narrative.scope_overview)

    _add_heading(doc, "4 Deliverables and acceptance criteria", level=1)
    for item in stored.narrative.deliverables:
        _add_heading(doc, item.title, level=2)
        _add_rich_paragraph(doc, item.description)
        for crit in item.acceptance_criteria:
            doc.add_paragraph(sanitize_docx_text(crit), style="List Bullet")

    _add_heading(doc, "6 Timeline and milestones", level=1)
    _add_table(
        doc,
        ["Phase", "Weeks", "Start", "End", "Indicative start", "Indicative end"],
        [
            [
                r.phase,
                str(r.planned_weeks),
                str(r.start_week),
                str(r.end_week),
                r.indicative_start or "",
                r.indicative_end or "",
            ]
            for r in inputs.timeline_table
        ],
    )

    _add_heading(doc, "14 Commercials", level=1)
    _add_rich_paragraph(doc, inputs.commercials.fee_display)
    _add_table(
        doc,
        ["Label", "Percent", "Amount", "Trigger"],
        [
            [
                str(row.get("label") or ""),
                str(row.get("percent") or ""),
                str(row.get("amount_display") or ""),
                str(row.get("trigger") or ""),
            ]
            for row in inputs.commercials.payment_rows
        ],
    )

    _add_heading(doc, "16 Signatures", level=1)
    _add_table(
        doc,
        ["Party", "Name", "Title", "Signature", "Date"],
        [
            ["Client", "", "", "", ""],
            ["Company", meta.created_by_name, "", "", ""],
        ],
    )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
