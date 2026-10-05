from __future__ import annotations

import csv
import email
from dataclasses import dataclass
from io import BytesIO, StringIO

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings
from app.services.file_validation import DetectedFileType

TRUNCATION_MARKER = "\n\n[TRUNCATED]"


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    truncated: bool
    empty_reason: str | None = None


def _apply_text_cap(raw: str) -> ExtractionResult:
    settings = get_settings()
    limit = settings.MAX_EXTRACTED_TEXT_CHARS
    if len(raw) <= limit:
        return ExtractionResult(text=raw, truncated=False)
    return ExtractionResult(
        text=raw[:limit] + TRUNCATION_MARKER,
        truncated=True,
    )


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> ExtractionResult:
    reader = PdfReader(BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    joined = "\n".join(parts).strip()
    if not joined:
        return ExtractionResult(text="", truncated=False, empty_reason="no_text_layer")
    return _apply_text_cap(joined)


def _extract_docx(content: bytes) -> ExtractionResult:
    doc = Document(BytesIO(content))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text:
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return _apply_text_cap("\n".join(parts))


def _extract_eml(content: bytes) -> ExtractionResult:
    msg = email.message_from_bytes(content)
    headers = [
        f"From: {msg.get('From', '')}",
        f"To: {msg.get('To', '')}",
        f"Subject: {msg.get('Subject', '')}",
        f"Date: {msg.get('Date', '')}",
    ]
    body = ""
    if msg.is_multipart():
        attachment_names: list[str] = []
        for part in msg.walk():
            disposition = part.get_content_disposition()
            if disposition == "attachment":
                filename = part.get_filename()
                if filename:
                    attachment_names.append(filename)
            elif part.get_content_type() == "text/plain" and not body:
                payload = part.get_payload(decode=True)
                if isinstance(payload, bytes):
                    body = _decode_text(payload)
        if attachment_names:
            headers.append("Attachments: " + ", ".join(attachment_names))
    else:
        payload = msg.get_payload(decode=True)
        if isinstance(payload, bytes):
            body = _decode_text(payload)
    text = "\n".join(headers) + "\n\n" + body
    return _apply_text_cap(text.strip())


def _extract_csv(content: bytes) -> ExtractionResult:
    decoded = _decode_text(content)
    reader = csv.reader(StringIO(decoded))
    rows = ["\t".join(row) for row in reader]
    return _apply_text_cap("\n".join(rows))


def _extract_xlsx(content: bytes) -> ExtractionResult:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                parts.append("\t".join(values))
    wb.close()
    return _apply_text_cap("\n".join(parts))


def extract_text(content: bytes, file_type: DetectedFileType) -> ExtractionResult:
    if file_type == DetectedFileType.PDF:
        return _extract_pdf(content)
    if file_type == DetectedFileType.DOCX:
        return _extract_docx(content)
    if file_type == DetectedFileType.EML:
        return _extract_eml(content)
    if file_type == DetectedFileType.CSV:
        return _extract_csv(content)
    if file_type == DetectedFileType.XLSX:
        return _extract_xlsx(content)
    if file_type in (DetectedFileType.TXT, DetectedFileType.MD):
        return _apply_text_cap(_decode_text(content))
    return _apply_text_cap(_decode_text(content))
